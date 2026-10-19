
import json
import urllib.request
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# --------------------------------------------------
# SBG-v4: Generate Chapter 1 workbook, all 47 verses
# --------------------------------------------------

BASE_URL = (
    "https://raw.githubusercontent.com/"
    "AdiyogiArts/gita-json/main/base.json"
)
EN_URL = (
    "https://raw.githubusercontent.com/"
    "AdiyogiArts/gita-json/main/en.json"
)

OUTPUT = Path(
    r"C:\SBG\SBG-v4\data"
    r"\Bhagavad_Gita_Chapter1_All_47_Verses.xlsx"
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)


def download_json(url):
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "SBG-v4-workbook-generator"}
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        return json.loads(response.read().decode("utf-8"))


def clean(value):
    if not isinstance(value, str):
        return ""
    return " ".join(value.replace("\r", "\n").split())


def make_hook(verse, explanation, translation):
    # Prefer a verse-specific idea from the explanation.
    source = clean(explanation) or clean(translation)
    if source:
        first_sentence = source.split(".")[0].strip()
        if first_sentence:
            return first_sentence[:150]
    return (
        f"What can Bhagavad Gita Chapter 1, "
        f"Verse {verse} teach us today?"
    )


def make_life_lesson(explanation, translation):
    explanation = clean(explanation)
    translation = clean(translation)

    if explanation:
        # Keep the reflection suitable for short-form narration.
        sentences = [
            part.strip()
            for part in explanation.replace("!", ".").split(".")
            if part.strip()
        ]
        reflection = ". ".join(sentences[:2]).strip()
        if reflection:
            return reflection.rstrip(".") + "."

    if translation:
        return (
            "Reflect on this moment before reacting. "
            "Consider its meaning, the people affected, "
            "and the consequences of your choices."
        )

    return (
        "Pause, seek clarity, and consider the consequences "
        "before taking action."
    )


def make_scene3(translation):
    translation = clean(translation)
    if not translation:
        return (
            "This verse is part of the dialogue on the battlefield "
            "of Kurukshetra. Review the translation before rendering."
        )
    return translation


def make_scene4(lesson):
    return lesson


print("Downloading Bhagavad Gita source data...")
base_data = download_json(BASE_URL)
english_data = download_json(EN_URL)

rows = []

for verse_number in range(1, 48):
    key = f"1-{verse_number}"

    if key not in base_data:
        raise ValueError(f"Sanskrit source is missing verse {key}.")

    if key not in english_data:
        raise ValueError(f"English source is missing verse {key}.")

    source = base_data[key]
    english = english_data[key]

    sanskrit = clean(source.get("sanskrit", ""))
    transliteration = clean(source.get("transliteration", ""))
    translation = clean(english.get("translation", ""))
    explanation = clean(english.get("explanation", ""))

    if not sanskrit or not translation:
        raise ValueError(
            f"Verse {key} has missing Sanskrit or English translation. "
            "Workbook was not created."
        )

    lesson = make_life_lesson(explanation, translation)

    if verse_number == 1:
        scene1 = (
            "We are reciting Shree Bhagavad Gita, "
            "Chapter 1, Shloka 1."
        )
        scene3 = (
            "Dhritarashtra asks Sanjaya what happened when "
            "the Kauravas and Pandavas gathered at Kurukshetra, "
            "ready for battle."
        )
        scene4 = (
            "Before reacting to conflict, pause and seek clarity. "
            "Sometimes, wisdom begins simply by choosing to see clearly."
        )
    else:
        scene1 = (
            "We are reciting Shree Bhagavad Gita, "
            f"Chapter 1, Shloka {verse_number}."
        )
        scene3 = make_scene3(translation)
        scene4 = make_scene4(lesson)

    hook = make_hook(verse_number, explanation, translation)

    rows.append({
        "Chapter": 1,
        "Verse": verse_number,
        "index": verse_number,
        "Sanskrit": sanskrit,
        "Transliteration": transliteration,
        "English Translation": translation,
        "Key Theme": "Chapter 1: Arjuna's Dilemma",
        "Life Lesson": lesson,
        "Hook": hook,
        "Narration Prompt": (
            f"Prepare a clear, respectful YouTube Shorts narration "
            f"for Bhagavad Gita Chapter 1, Verse {verse_number}. "
            "Explain the verse accurately and offer a practical "
            "reflection. Avoid inventing scriptural claims."
        ),
        "Generated Narration": f"{scene3} {scene4}",
        "Voice Style": (
            "Scene 2: Male; Scenes 1, 3, 4: Female"
        ),
        "Image Prompt": (
            f"Vertical 9:16 devotional cinematic illustration for "
            f"Bhagavad Gita Chapter 1 Verse {verse_number}. "
            "Respect the verse context, maintain consistent character "
            "designs, and use no text or watermark."
        ),
        "Video Prompt": (
            "Animate the approved scene image with subtle cinematic "
            "movement, natural cloth and environmental motion. "
            "Respectful devotional tone, vertical 1080x1920, "
            "no added text."
        ),
        "Subtitle Text": sanskrit,
        "Image File": "",
        "Audio File": "",
        "Subtitle File": "",
        "Video File": "",
        "Scene 1 Narration": scene1,
        "Scene 2 Narration": sanskrit,
        "Scene 3 Narration": scene3,
        "Scene 4 Narration": scene4,
        "Upload Status": "Pending",
    })


# Put Chapter and Verse first because some workbook readers
# identify the target row using the first two columns.
columns = list(rows[0].keys())

wb = Workbook()
ws = wb.active
ws.title = "Content_Master_V1-47"

ws.append(columns)

for row in rows:
    ws.append([row[column] for column in columns])

# Format the content sheet.
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="17365D")
    cell.alignment = Alignment(
        wrap_text=True,
        vertical="top"
    )

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(
            wrap_text=True,
            vertical="top"
        )

for col_index, column_name in enumerate(columns, start=1):
    ws.column_dimensions[get_column_letter(col_index)].width = 22

wide_columns = [
    "Sanskrit",
    "Transliteration",
    "English Translation",
    "Life Lesson",
    "Hook",
    "Narration Prompt",
    "Generated Narration",
    "Image Prompt",
    "Video Prompt",
    "Scene 1 Narration",
    "Scene 2 Narration",
    "Scene 3 Narration",
    "Scene 4 Narration",
]

for column_name in wide_columns:
    col_index = columns.index(column_name) + 1
    ws.column_dimensions[get_column_letter(col_index)].width = 48

for row_index in range(2, ws.max_row + 1):
    ws.row_dimensions[row_index].height = 72


# Add workbook instructions.
readme = wb.create_sheet("Read_Me")
readme.append(["Setting", "Details"])
notes = [
    (
        "Coverage",
        "Chapter 1, verses 1-47; 47 rows."
    ),
    (
        "Scene 1",
        "Introduces the chapter and verse."
    ),
    (
        "Scene 2",
        "Sanskrit recitation; male voice per SBG-v4 preference."
    ),
    (
        "Scene 3",
        "English explanation/translation; female voice."
    ),
    (
        "Scene 4",
        "Reflection or practical lesson; female voice."
    ),
    (
        "Verse 1",
        "Preserves the approved Scene 1, Scene 3 and Scene 4 "
        "narration text supplied for the current project. "
        "Scene 2 uses the downloaded Sanskrit text."
    ),
    (
        "Important",
        "Translations and explanations are sourced from a public "
        "dataset and may need editorial review. Check Sanskrit "
        "against your preferred edition before publishing."
    ),
    (
        "Pipeline",
        "Update modules/sheet.py to use this workbook and the "
        "Content_Master_V1-47 sheet. Then update the pipeline's "
        "verse-selection logic to process verses 1-47."
    ),
    (
        "Output",
        str(OUTPUT)
    ),
]

for note in notes:
    readme.append(list(note))

readme.freeze_panes = "A2"
readme.column_dimensions["A"].width = 24
readme.column_dimensions["B"].width = 110

for cell in readme[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="17365D")

for row in readme.iter_rows():
    for cell in row:
        cell.alignment = Alignment(
            wrap_text=True,
            vertical="top"
        )

# Validate before saving.
required = {
    "Chapter",
    "Verse",
    "Sanskrit",
    "English Translation",
    "Life Lesson",
    "Hook",
    "Scene 1 Narration",
    "Scene 2 Narration",
    "Scene 3 Narration",
    "Scene 4 Narration",
    "Upload Status",
}

missing = required - set(columns)
if missing:
    raise ValueError(f"Required columns missing: {sorted(missing)}")

for row in rows:
    for scene in (
        "Scene 1 Narration",
        "Scene 2 Narration",
        "Scene 3 Narration",
        "Scene 4 Narration",
    ):
        if not clean(row[scene]):
            raise ValueError(
                f"Verse {row['Verse']} has empty {scene}."
            )

wb.save(OUTPUT)

# Reopen to verify the saved workbook.
check = load_workbook(OUTPUT, read_only=True, data_only=True)
check_ws = check["Content_Master_V1-47"]

assert check_ws.max_row == 48, (
    f"Expected 47 verses plus header; got {check_ws.max_row} rows."
)

print()
print("WORKBOOK CREATED SUCCESSFULLY")
print(f"File: {OUTPUT}")
print("Verses: 47")
print("Scene narration fields: 4 per verse")
print("Total scene narration entries: 188")
print("Sheet: Content_Master_V1-47")
print()
print("Next: update modules/sheet.py and pipeline verse iteration.")
