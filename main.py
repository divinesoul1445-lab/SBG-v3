from pathlib import Path
import shutil
import re

from modules.pipeline import Pipeline



import os
import logging
import warnings

# Suppress routine library logging
os.environ["TF_CPP_MIN_LOG_LEVEL"] = "3"
os.environ["HF_HUB_DISABLE_PROGRESS_BARS"] = "1"
os.environ["TOKENIZERS_PARALLELISM"] = "false"

logging.disable(logging.WARNING)
warnings.filterwarnings("ignore")


def clear_output_folder():
    output_folder = Path("output")

    if output_folder.exists():
        print()
        print("=" * 80)
        print("CLEANING PREVIOUS OUTPUT")
        print("=" * 80)
        print(f"Removing: {output_folder.resolve()}")
        shutil.rmtree(output_folder)
        print("Previous output folder removed.")

    output_folder.mkdir(parents=True, exist_ok=True)
    print(f"Fresh output folder created: {output_folder.resolve()}")
    print()


def export_final_videos():
    output_folder = Path("output")
    destination = Path("final_videos")
    destination.mkdir(parents=True, exist_ok=True)

    copied = 0

    for video in output_folder.glob(
        "chapter_*/verse_*/final/merged_video.mp4"
    ):
        match = re.fullmatch(
            r"chapter_(\d+)/verse_(\d+)",
            f"{video.parents[2].name}/{video.parents[1].name}",
        )

        if not match:
            continue

        chapter, verse = map(int, match.groups())
        target = destination / (
            f"Chapter_{chapter:03d}Verse_{verse:03d}.mp4"
        )

        shutil.copy2(video, target)
        print(f"Exported: {target}")

        copied += 1

    print()
    print(f"Export complete: {copied} video(s) in {destination.resolve()}")



def ensure_tracking_columns(excel_manager):
    """Add render/export tracking columns without changing Upload Status."""
    from openpyxl import load_workbook

    workbook_path = Path(excel_manager.excel_file)
    workbook = load_workbook(workbook_path)
    try:
        worksheet = workbook[excel_manager.SHEET_NAME]
        headers = {
            str(cell.value).strip(): cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }
        for header in (
            "Render Status",
            "Local Export Status",
            "Google Drive Status",
        ):
            if header not in headers:
                worksheet.cell(1, worksheet.max_column + 1).value = header
        workbook.save(workbook_path)
    finally:
        workbook.close()


def update_tracking_status(excel_manager, chapter, verse, **updates):
    """Write status fields for one verse to the content workbook."""
    from openpyxl import load_workbook

    workbook_path = Path(excel_manager.excel_file)
    workbook = load_workbook(workbook_path)
    try:
        worksheet = workbook[excel_manager.SHEET_NAME]
        headers = {
            str(cell.value).strip(): cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }
        for row in range(2, worksheet.max_row + 1):
            row_chapter = excel_manager._to_int(
                worksheet.cell(row, headers["Chapter"]).value
            )
            row_verse = excel_manager._to_int(
                worksheet.cell(row, headers["Verse"]).value
            )
            if row_chapter == chapter and row_verse == verse:
                for header, value in updates.items():
                    worksheet.cell(row, headers[header]).value = value
                workbook.save(workbook_path)
                return
        raise ValueError(f"Chapter {chapter}, Verse {verse} not found in workbook")
    finally:
        workbook.close()


def get_tracking_status(excel_manager, chapter, verse):
    """Read the current status of a verse from Excel."""
    from openpyxl import load_workbook

    workbook = load_workbook(excel_manager.excel_file, read_only=True)
    try:
        worksheet = workbook[excel_manager.SHEET_NAME]
        headers = {
            str(cell.value).strip(): cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }

        for row in range(2, worksheet.max_row + 1):
            ch = excel_manager._to_int(
                worksheet.cell(row, headers["Chapter"]).value
            )
            vs = excel_manager._to_int(
                worksheet.cell(row, headers["Verse"]).value
            )

            if ch == chapter and vs == verse:
                return {
                    name: str(worksheet.cell(row, col).value or "").strip()
                    for name, col in headers.items()
                    if name in (
                        "Render Status",
                        "Local Export Status",
                        "Google Drive Status",
                    )
                }

        raise ValueError(
            f"Chapter {chapter}, Verse {verse} not found in Excel"
        )
    finally:
        workbook.close()


def main():
    import shutil

    pipeline = Pipeline()
    ensure_tracking_columns(pipeline.excel)

    failures = []
    successes = []

    destinations = [
        Path("final_videos"),
        Path(r"G:\My Drive\Final Verses"),
    ]

    for destination in destinations:
        destination.mkdir(parents=True, exist_ok=True)

    for verse_number in range(1, 11):
        chapter = 1
        target_name = f"Chapter_{chapter:03d}Verse_{verse_number:03d}.mp4"

        output_video = (
            Path("output")
            / "chapter_001"
            / f"verse_{verse_number:03d}"
            / "final"
            / "merged_video.mp4"
        )

        local_video = destinations[0] / target_name
        drive_video = destinations[1] / target_name

        print("\n" + "=" * 60)
        print(f"CHAPTER {chapter}, VERSE {verse_number}")
        print("=" * 60)

        try:
            status = get_tracking_status(
                pipeline.excel, chapter, verse_number
            )

            # Prefer the original rendered output, then existing exports.
            candidates = [output_video, local_video, drive_video]

            source_video = next(
                (
                    path for path in candidates
                    if path.is_file() and path.stat().st_size > 0
                ),
                None,
            )

            render_completed = (
                status.get("Render Status", "").casefold()
                == "completed"
            )

            if render_completed and source_video is not None:
                print(f"SKIPPING RENDER: Existing video found: {source_video}")

            elif source_video == output_video and output_video.is_file():
                # Recover a video already rendered before a previous
                # interruption, even if Excel was not updated.
                print("Existing rendered video found. Skipping rerender.")

            else:
                # No usable existing video: rendering is necessary.
                print("No usable completed video found. Rendering now.")

                update_tracking_status(
                    pipeline.excel, chapter, verse_number,
                    **{
                        "Render Status": "Processing",
                        "Local Export Status": "Pending",
                        "Google Drive Status": "Pending",
                    },
                )

                success = pipeline.run(
                    chapter=chapter, verse=verse_number
                )

                if not success:
                    update_tracking_status(
                        pipeline.excel, chapter, verse_number,
                        **{"Render Status": "Failed"},
                    )
                    failures.append(
                        (verse_number, "Pipeline returned False")
                    )
                    continue

                if (
                    not output_video.is_file()
                    or output_video.stat().st_size <= 0
                ):
                    update_tracking_status(
                        pipeline.excel, chapter, verse_number,
                        **{"Render Status": "Failed"},
                    )
                    failures.append(
                        (verse_number, "Rendered video missing or empty")
                    )
                    continue

                source_video = output_video

            # A valid video exists. Update status and recover exports
            # independently, without rerendering.
            update_tracking_status(
                pipeline.excel, chapter, verse_number,
                **{"Render Status": "Completed"},
            )

            copy_errors = []

            for destination in destinations:
                target = destination / target_name

                try:
                    # Avoid copying a file onto itself.
                    same_file = (
                        source_video.resolve() == target.resolve()
                    )

                    if not same_file:
                        shutil.copy2(source_video, target)

                    if (
                        not target.is_file()
                        or target.stat().st_size <= 0
                        or target.stat().st_size
                        != source_video.stat().st_size
                    ):
                        raise IOError("Video copy verification failed")

                    header = (
                        "Local Export Status"
                        if destination == destinations[0]
                        else "Google Drive Status"
                    )

                    update_tracking_status(
                        pipeline.excel, chapter, verse_number,
                        **{header: "Completed"},
                    )

                    print(f"VERIFIED: {target}")

                except Exception as exc:
                    header = (
                        "Local Export Status"
                        if destination == destinations[0]
                        else "Google Drive Status"
                    )

                    update_tracking_status(
                        pipeline.excel, chapter, verse_number,
                        **{header: "Failed"},
                    )

                    copy_errors.append(f"{target}: {exc}")
                    print(f"EXPORT FAILED: {target}: {exc}")

            if copy_errors:
                failures.append((verse_number, "; ".join(copy_errors)))
            else:
                successes.append(verse_number)
                print(f"SUCCESS: Verse {verse_number}")

        except Exception as exc:
            failures.append((verse_number, str(exc)))
            print(f"FAILED: Verse {verse_number}: {exc}")

    print("\n" + "=" * 60)
    print(f"Successful verses: {successes}")
    print(f"Failures: {len(failures)}")

    for verse_number, error in failures:
        print(f"Verse {verse_number}: {error}")


if __name__ == "__main__":
    import time

    start_time = time.perf_counter()

    try:
        main()
    finally:
        elapsed = time.perf_counter() - start_time
        hours, remainder = divmod(int(elapsed), 3600)
        minutes, seconds = divmod(remainder, 60)

        print()
        print("=" * 60)
        print(
            f"TOTAL PIPELINE TIME: "
            f"{hours:02d}:{minutes:02d}:{seconds:02d}"
        )
        print("=" * 60)

