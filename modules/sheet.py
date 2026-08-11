"""
Excel Content Manager - SBG V3

The Excel Content Master is the approved source of truth
for production content.

Current development mode:
Chapter 1, Verse 1 only.

Excel structure:

```
index
Chapter
Verse
Sanskrit
English Translation
Life Lesson
Hook
Scene 1 Narration
Scene 2 Narration
Scene 3 Narration
Scene 4 Narration
Upload Status
```

"""

# from future import annotations

from pathlib import Path

from openpyxl import load_workbook

# ============================================================

# EXCEL FILE

# ============================================================

EXCEL_FILE = "C:\\SBG\\SBG-v3\\data\\SBG_V3_Content_Master_Chapter1_Verse1.xlsx"

# ============================================================

# EXCEL CONTENT MANAGER

# ============================================================

class ExcelContentManager:


    SHEET_NAME = "Content_Master"

    # --------------------------------------------------------
    # DEVELOPMENT LOCK
    #
    # For now, the entire pipeline is restricted to:
    #
    # Chapter 1
    # Verse 1
    # --------------------------------------------------------

    ACTIVE_CHAPTER = 1
    ACTIVE_VERSE = 1

    # --------------------------------------------------------
    # Expected Excel columns
    # --------------------------------------------------------

    REQUIRED_COLUMNS = [
        "index",
        "Chapter",
        "Verse",
        "Sanskrit",
        "English Translation",
        "Life Lesson",
        "Hook",
        "Scene 1 Narration",
        "Scene 2 Narration",
        "Scene 3 Narration",
        "Scene 4 Narration",
        "Upload Status",
    ]

    # ========================================================
    # INIT
    # ========================================================

    def __init__(
        self,
        excel_file: Path | str = EXCEL_FILE,
    ):

        self.excel_file = Path(excel_file)

        if not self.excel_file.exists():

            raise FileNotFoundError(
                "Excel content master not found:\n"
                f"{self.excel_file}"
            )

    # ========================================================
    # LOAD WORKSHEET
    # ========================================================

    def _load_sheet(self):

        workbook = load_workbook(
            filename=self.excel_file,
            read_only=True,
            data_only=True,
        )

        if self.SHEET_NAME not in workbook.sheetnames:

            available_sheets = workbook.sheetnames

            workbook.close()

            raise Exception(
                f"Worksheet '{self.SHEET_NAME}' "
                "not found.\n"
                f"Available worksheets: "
                f"{available_sheets}"
            )

        worksheet = workbook[self.SHEET_NAME]

        return workbook, worksheet

    # ========================================================
    # READ EXCEL ROWS
    # ========================================================

    def _rows(self):

        workbook, worksheet = self._load_sheet()

        try:

            rows = worksheet.iter_rows(
                values_only=True
            )

            # ------------------------------------------------
            # Header row
            # ------------------------------------------------

            try:

                raw_headers = next(rows)

            except StopIteration:

                raise Exception(
                    "Excel worksheet is empty."
                )

            headers = [
                str(header).strip()
                if header is not None
                else ""
                for header in raw_headers
            ]

            # ------------------------------------------------
            # Validate structure
            # ------------------------------------------------

            self._validate_columns(
                headers
            )

            # ------------------------------------------------
            # Convert rows to dictionaries
            # ------------------------------------------------

            for row in rows:

                # Skip completely empty rows.

                if not any(
                    value is not None
                    for value in row
                ):
                    continue

                data = {}

                for index, header in enumerate(
                    headers
                ):

                    if not header:
                        continue

                    value = (
                        row[index]
                        if index < len(row)
                        else None
                    )

                    data[header] = value

                yield data

        finally:

            workbook.close()

    # ========================================================
    # VALIDATE EXCEL COLUMNS
    # ========================================================

    def _validate_columns(
        self,
        headers: list[str],
    ):

        missing_columns = [
            column
            for column in self.REQUIRED_COLUMNS
            if column not in headers
        ]

        if missing_columns:

            raise Exception(
                "Excel content master has missing "
                "required columns:\n"
                + "\n".join(
                    f"  - {column}"
                    for column in missing_columns
                )
            )

    # ========================================================
    # GET VERSE
    # ========================================================

    def get_verse(
        self,
        chapter: int,
        verse: int,
    ) -> dict:

        # ----------------------------------------------------
        # Development safety lock
        #
        # Nothing except Chapter 1 / Verse 1 is allowed.
        # ----------------------------------------------------

        if (
            chapter != self.ACTIVE_CHAPTER
            or verse != self.ACTIVE_VERSE
        ):

            raise Exception(
                "SBG V3 is currently locked to "
                "Chapter 1, Verse 1 only."
            )

        # ----------------------------------------------------
        # Search Excel
        # ----------------------------------------------------

        for row in self._rows():

            row_chapter = self._to_int(
                row.get("Chapter")
            )

            row_verse = self._to_int(
                row.get("Verse")
            )

            if (
                row_chapter == chapter
                and row_verse == verse
            ):

                return self._prepare_row(
                    row
                )

        # ----------------------------------------------------
        # Verse not found
        # ----------------------------------------------------

        raise Exception(
            f"Chapter {chapter}, Verse {verse} "
            "not found in Excel content master."
        )

    # ========================================================
    # PREPARE ROW
    # ========================================================

    def _prepare_row(
        self,
        row: dict,
    ) -> dict:

        result = {}

        # ----------------------------------------------------
        # Copy all required fields
        # ----------------------------------------------------

        for column in self.REQUIRED_COLUMNS:

            result[column] = self._clean(
                row.get(column)
            )

        # ----------------------------------------------------
        # Numeric fields
        # ----------------------------------------------------

        result["index"] = self._to_int(
            row.get("index")
        )

        result["Chapter"] = self._to_int(
            row.get("Chapter")
        )

        result["Verse"] = self._to_int(
            row.get("Verse")
        )

        return result

    # ========================================================
    # GET SCENE NARRATIONS
    # ========================================================

    def get_scene_narrations(
        self,
        chapter: int = ACTIVE_CHAPTER,
        verse: int = ACTIVE_VERSE,
    ) -> list[str]:

        content = self.get_verse(
            chapter,
            verse,
        )

        scenes = [
            content["Scene 1 Narration"],
            content["Scene 2 Narration"],
            content["Scene 3 Narration"],
            content["Scene 4 Narration"],
        ]

        # ----------------------------------------------------
        # Every scene must contain approved narration.
        # ----------------------------------------------------

        for index, narration in enumerate(
            scenes,
            start=1,
        ):

            if not narration:

                raise Exception(
                    f"Scene {index} Narration is empty "
                    f"for Chapter {chapter}, "
                    f"Verse {verse}."
                )

        return scenes

    # ========================================================
    # GET FULL NARRATION
    # ========================================================

    def get_full_narration(
        self,
        chapter: int = ACTIVE_CHAPTER,
        verse: int = ACTIVE_VERSE,
    ) -> str:

        scenes = self.get_scene_narrations(
            chapter,
            verse,
        )

        return " ".join(
            scene.strip()
            for scene in scenes
            if scene.strip()
        )

    # ========================================================
    # UPDATE UPLOAD STATUS
    # ========================================================

    def update_upload_status(
        self,
        chapter: int,
        verse: int,
        status: str,
    ):

        # ----------------------------------------------------
        # Development safety lock
        # ----------------------------------------------------

        if (
            chapter != self.ACTIVE_CHAPTER
            or verse != self.ACTIVE_VERSE
        ):

            raise Exception(
                "SBG V3 is currently locked to "
                "Chapter 1, Verse 1 only."
            )

        workbook = load_workbook(
            filename=self.excel_file
        )

        try:

            # ------------------------------------------------
            # Verify worksheet
            # ------------------------------------------------

            if self.SHEET_NAME not in (
                workbook.sheetnames
            ):

                raise Exception(
                    f"Worksheet '{self.SHEET_NAME}' "
                    "not found."
                )

            worksheet = workbook[
                self.SHEET_NAME
            ]

            # ------------------------------------------------
            # Build header map
            # ------------------------------------------------

            headers = {}

            for cell in worksheet[1]:

                if cell.value is None:
                    continue

                header = str(
                    cell.value
                ).strip()

                headers[header] = cell.column

            # ------------------------------------------------
            # Required columns
            # ------------------------------------------------

            required = {
                "Chapter",
                "Verse",
                "Upload Status",
            }

            missing = (
                required
                - set(headers.keys())
            )

            if missing:

                raise Exception(
                    "Cannot update Upload Status. "
                    "Missing columns: "
                    + ", ".join(
                        sorted(missing)
                    )
                )

            # ------------------------------------------------
            # Find target row
            # ------------------------------------------------

            for row_number in range(
                2,
                worksheet.max_row + 1,
            ):

                row_chapter = self._to_int(
                    worksheet.cell(
                        row_number,
                        headers["Chapter"],
                    ).value
                )

                row_verse = self._to_int(
                    worksheet.cell(
                        row_number,
                        headers["Verse"],
                    ).value
                )

                if (
                    row_chapter == chapter
                    and row_verse == verse
                ):

                    worksheet.cell(
                        row_number,
                        headers["Upload Status"],
                    ).value = status

                    workbook.save(
                        self.excel_file
                    )

                    return

            raise Exception(
                f"Chapter {chapter}, Verse {verse} "
                "not found while updating "
                "Upload Status."
            )

        finally:

            workbook.close()

    # ========================================================
    # HELPERS
    # ========================================================

    @staticmethod
    def _clean(
        value,
    ) -> str:

        if value is None:
            return ""

        return str(value).strip()

    @staticmethod
    def _to_int(
        value,
    ):

        if value is None:
            return None

        try:

            return int(value)

        except (
            TypeError,
            ValueError,
        ):

            return None
